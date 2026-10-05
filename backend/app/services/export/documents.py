"""Експорт тесту в DOCX і PDF: для студентів (без відповідей) або з відповідями та ключем.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from html import escape
from pathlib import Path

LETTERS = "АБВГДЕЖЗИК"
TYPE_HINT = {
    "single_choice": "Оберіть одну правильну відповідь.",
    "multiple_choice": "Оберіть усі правильні відповіді.",
    "true_false": "Визначте, чи є твердження правдивим.",
}

# DejaVu — у Docker-образі, Arial — для запуску тестів на Windows
FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
]


@dataclass(frozen=True)
class ExportOption:
    text: str
    is_correct: bool


@dataclass(frozen=True)
class ExportQuestion:
    type: str
    text: str
    options: list[ExportOption]
    explanation: str | None = None


@dataclass(frozen=True)
class ExportTest:
    title: str
    questions: list[ExportQuestion]
    description: str | None = None
    subtitle: str | None = None          # наприклад, назви матеріалів


def correct_letters(q: ExportQuestion) -> str:
    return ", ".join(LETTERS[i] for i, o in enumerate(q.options) if o.is_correct)


# --- DOCX ---
def build_docx(test: ExportTest, with_answers: bool) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_heading(test.title, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if with_answers:
        p = doc.add_paragraph("Варіант для викладача — з відповідями")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True
    if test.subtitle:
        p = doc.add_paragraph(test.subtitle)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.color.rgb = RGBColor(0x6B, 0x70, 0x86)
    if test.description:
        doc.add_paragraph(test.description)

    if not with_answers:
        doc.add_paragraph("ПІБ: ____________________________    Група: __________    Дата: __________")
    doc.add_paragraph(f"Кількість питань: {len(test.questions)}")

    for n, q in enumerate(test.questions, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"{n}. ").bold = True
        p.add_run(q.text).bold = True
        hint = doc.add_paragraph(TYPE_HINT.get(q.type, ""))
        hint.paragraph_format.keep_with_next = True
        hint.runs[0].italic = True
        hint.runs[0].font.size = Pt(9)

        for i, o in enumerate(q.options):
            mark = "☐"
            correct = with_answers and o.is_correct
            if with_answers:
                mark = "☑" if o.is_correct else "☐"
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Pt(18)
            op.paragraph_format.space_after = Pt(0)
            op.paragraph_format.keep_with_next = i < len(q.options) - 1
            run = op.add_run(f"{mark} {LETTERS[i]}) {o.text}")
            if correct:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x9D, 0x55)

        if with_answers and q.explanation:
            ex = doc.add_paragraph()
            ex.paragraph_format.left_indent = Pt(18)
            r = ex.add_run(f"Пояснення: {q.explanation}")
            r.italic = True
            r.font.size = Pt(9)

    if with_answers and test.questions:
        doc.add_page_break()
        doc.add_heading("Ключ відповідей", level=2)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "Питання"
        table.rows[0].cells[1].text = "Правильна відповідь"
        for n, q in enumerate(test.questions, start=1):
            row = table.add_row().cells
            row[0].text = str(n)
            row[1].text = correct_letters(q)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- PDF ---
_fonts_registered: tuple[str, str] | None = None


def _register_fonts() -> tuple[str, str]:
    """Реєструє TTF-шрифт із кирилицею для reportlab."""
    global _fonts_registered
    if _fonts_registered:
        return _fonts_registered
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold in FONT_CANDIDATES:
        if Path(regular).exists() and Path(bold).exists():
            pdfmetrics.registerFont(TTFont("ExportSans", regular))
            pdfmetrics.registerFont(TTFont("ExportSans-Bold", bold))
            _fonts_registered = ("ExportSans", "ExportSans-Bold")
            return _fonts_registered
    raise RuntimeError("Не знайдено шрифт із кирилицею для PDF (очікувався DejaVu Sans)")


def build_pdf(test: ExportTest, with_answers: bool) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    regular, bold = _register_fonts()
    base = ParagraphStyle("base", fontName=regular, fontSize=10.5, leading=14)
    st = {
        "title": ParagraphStyle("title", parent=base, fontName=bold, fontSize=16, leading=20, alignment=TA_CENTER),
        "center": ParagraphStyle("center", parent=base, alignment=TA_CENTER, textColor=colors.HexColor("#6b7086")),
        "q": ParagraphStyle("q", parent=base, fontName=bold, spaceBefore=8),
        "hint": ParagraphStyle("hint", parent=base, fontSize=8.5, textColor=colors.HexColor("#6b7086")),
        "opt": ParagraphStyle("opt", parent=base, leftIndent=14),
        "opt_ok": ParagraphStyle("opt_ok", parent=base, leftIndent=14, fontName=bold, textColor=colors.HexColor("#1f9d55")),
        "expl": ParagraphStyle("expl", parent=base, leftIndent=14, fontSize=9, textColor=colors.HexColor("#444444")),
        "h2": ParagraphStyle("h2", parent=base, fontName=bold, fontSize=13, leading=17, spaceAfter=6),
    }

    story = [Paragraph(escape(test.title), st["title"]), Spacer(1, 3 * mm)]
    if with_answers:
        story.append(Paragraph("Варіант для викладача — з відповідями", st["center"]))
    if test.subtitle:
        story.append(Paragraph(escape(test.subtitle), st["center"]))
    if test.description:
        story += [Spacer(1, 2 * mm), Paragraph(escape(test.description), base)]
    story.append(Spacer(1, 4 * mm))
    if not with_answers:
        story += [Paragraph("ПІБ: ______________________________   Група: ___________   Дата: ___________", base),
                  Spacer(1, 2 * mm)]
    story.append(Paragraph(f"Кількість питань: {len(test.questions)}", base))

    for n, q in enumerate(test.questions, start=1):
        block = [
            Paragraph(f"{n}. {escape(q.text)}", st["q"]),
            Paragraph(escape(TYPE_HINT.get(q.type, "")), st["hint"]),
        ]
        for i, o in enumerate(q.options):
            ok = with_answers and o.is_correct
            mark = "☑" if ok else "☐"
            block.append(Paragraph(f"{mark} {LETTERS[i]}) {escape(o.text)}", st["opt_ok" if ok else "opt"]))
        if with_answers and q.explanation:
            block.append(Paragraph(f"Пояснення: {escape(q.explanation)}", st["expl"]))
        story.append(KeepTogether(block))

    if with_answers and test.questions:
        story += [PageBreak(), Paragraph("Ключ відповідей", st["h2"])]
        rows = [["Питання", "Правильна відповідь"]] + [[str(n), correct_letters(q)]
                                                       for n, q in enumerate(test.questions, start=1)]
        table = Table(rows, colWidths=[30 * mm, 60 * mm])
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), regular),
            ("FONTNAME", (0, 0), (-1, 0), bold),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c9cbe0")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeecfd")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ]))
        story.append(table)

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                      topMargin=16 * mm, bottomMargin=16 * mm, title=test.title).build(story)
    return buf.getvalue()
