"""Витягування тексту з PDF/DOCX у документи LangChain з метаданими (сторінка/розділ)."""
from __future__ import annotations

import re
from pathlib import Path

from langchain_core.documents import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class UnsupportedFileError(ValueError):
    pass


MIN_TEXT_CHARS = 300


class EmptyDocumentError(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)      # перенесення слів через дефіс у PDF
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_pdf(path: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs: list[Document] = []
    for page_num, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:
            docs.append(Document(page_content=text, metadata={"page": page_num}))
    return docs


def load_docx(path: Path) -> list[Document]:
    """Розбиває DOCX за заголовками: кожен розділ — окремий документ."""
    import docx

    document = docx.Document(str(path))
    docs: list[Document] = []
    section_title = ""
    buffer: list[str] = []

    def flush() -> None:
        text = _clean("\n".join(buffer))
        if text:
            docs.append(Document(page_content=text, metadata={"section": section_title}))
        buffer.clear()

    for para in document.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        text = para.text.strip()
        if not text:
            continue
        if style.startswith("heading") or style.startswith("заголовок") or style == "title":
            flush()
            section_title = text[:200]
            buffer.append(text)
        else:
            buffer.append(text)
    flush()

    # Таблиці: рядок -> "клітинка | клітинка"
    for table in document.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        text = _clean("\n".join(r for r in rows if r.strip(" |")))
        if text:
            docs.append(Document(page_content=text, metadata={"section": "table"}))
    return docs


def load_document(path: str | Path) -> list[Document]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in (".pdf", ".docx"):
        raise UnsupportedFileError(f"Непідтримуваний формат: {ext}")
    try:
        docs = load_pdf(path) if ext == ".pdf" else load_docx(path)
    except Exception as exc:  # noqa: BLE001
        raise EmptyDocumentError(
            f"Не вдалося прочитати файл: він пошкоджений або захищений паролем ({type(exc).__name__})."
        ) from exc

    total = sum(len(d.page_content) for d in docs)
    if total == 0:
        raise EmptyDocumentError(
            "Не вдалося витягти текст. Можливо, це скан без текстового шару — потрібне розпізнавання (OCR)."
        )
    if total < MIN_TEXT_CHARS:
        raise EmptyDocumentError(
            f"У файлі замало тексту ({total} символів). Для генерації питань потрібен навчальний матеріал "
            f"обсягом щонайменше {MIN_TEXT_CHARS} символів (кілька абзаців)."
        )
    return docs
