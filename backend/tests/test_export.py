import io
import unittest

from app.services.export.documents import (
    ExportOption, ExportQuestion, ExportTest, build_docx, build_pdf, correct_letters,
)

TEST = ExportTest(
    title="Веб-фреймворки Go",
    subtitle="Матеріали: Лекція 3",
    questions=[
        ExportQuestion("single_choice", "Що таке горутина?", [
            ExportOption("Системний потік", False), ExportOption("Легковажний потік", True),
            ExportOption("Тип каналу", False), ExportOption("Пакет", False)],
            explanation="Горутиною керує рантайм Go."),
        ExportQuestion("true_false", "Go компілюється в машинний код.", [
            ExportOption("Правда", True), ExportOption("Неправда", False)]),
    ],
)


def docx_text(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        parts += [c.text for row in t.rows for c in row.cells]
    return "\n".join(parts)


class ExportTest_(unittest.TestCase):
    def test_correct_letters_ukrainian(self):
        self.assertEqual(correct_letters(TEST.questions[0]), "Б")
        self.assertEqual(correct_letters(TEST.questions[1]), "А")

    def test_docx_student_has_no_answers(self):
        text = docx_text(build_docx(TEST, with_answers=False))
        self.assertIn("Що таке горутина?", text)
        self.assertIn("ПІБ", text)
        self.assertNotIn("☑", text)
        self.assertNotIn("Ключ відповідей", text)
        self.assertNotIn("Горутиною керує", text)

    def test_docx_teacher_has_key_and_explanations(self):
        text = docx_text(build_docx(TEST, with_answers=True))
        self.assertIn("☑ Б) Легковажний потік", text)
        self.assertIn("Горутиною керує рантайм Go.", text)
        self.assertIn("Ключ відповідей", text)
        self.assertNotIn("ПІБ", text)

    def test_pdf_is_valid_and_keeps_cyrillic(self):
        data = build_pdf(TEST, with_answers=True)
        self.assertTrue(data.startswith(b"%PDF"))
        try:
            from pypdf import PdfReader
        except ImportError:
            self.skipTest("pypdf не установлен")
        text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)
        self.assertIn("горутина", text)
        self.assertIn("Ключ відповідей", text)

    def test_pdf_student_variant(self):
        self.assertTrue(build_pdf(TEST, with_answers=False).startswith(b"%PDF"))


if __name__ == "__main__":
    unittest.main()
